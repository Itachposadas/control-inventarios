# backend/reporte_excel.py
# Reporte de mantenimiento en Excel (.xlsx): una hoja de resumen y una por tabla.
# Las cantidades van como números (con formato de moneda) para poder sumarlas o
# filtrarlas en Excel; los totales son fórmulas.
from datetime import timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from models import utcnow

MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
TIPOS = {"preventivo": "Preventivo", "correctivo": "Correctivo", "emergencia": "Emergencia"}

MONEDA = '"$"#,##0.00'
PORCENTAJE = "0.0%"
FUENTE = "Calibri"
GRIS = PatternFill("solid", fgColor="E6E6E6")
FINA = Side(style="thin", color="BFBFBF")
GRUESA = Side(style="medium", color="000000")


def _fecha(iso):
    a, m, d = iso.split("-")
    return f"{d}/{m}/{a}"


def _mes(aaaamm):
    a, m = aaaamm.split("-")
    return f"{MESES[int(m) - 1]} {a}"


def _encabezado(ws, titulo, periodo, columnas):
    """Título, periodo y área arriba de cada hoja. Devuelve la fila donde va la tabla."""
    ws.append(["MUNICIPIO DE ATLACOMULCO · COORDINACIÓN DE PARQUE VEHICULAR"])
    ws.append([titulo.upper()])
    ws.append([f"Periodo: {_fecha(periodo['desde'])} al {_fecha(periodo['hasta'])}   ·   "
               f"Área: {periodo['area'] or 'Todas las áreas'}"])
    ws["A1"].font = Font(name=FUENTE, size=9, color="555555")
    ws["A2"].font = Font(name=FUENTE, size=13, bold=True)
    ws["A3"].font = Font(name=FUENTE, size=10, color="333333")
    for fila in (1, 2, 3):
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=max(columnas, 2))
    ws.append([])
    return 5


def _tabla(ws, fila, encabezados, filas, formatos, anchos, total=None):
    """Escribe una tabla con encabezado gris, filetes y (opcional) renglón de total."""
    for col, texto in enumerate(encabezados, start=1):
        c = ws.cell(row=fila, column=col, value=texto)
        c.font = Font(name=FUENTE, bold=True, size=10)
        c.fill = GRIS
        c.border = Border(top=GRUESA, bottom=GRUESA)
        c.alignment = Alignment(horizontal="right" if formatos[col - 1] else "left", vertical="center")
    inicio = fila + 1
    for i, datos in enumerate(filas):
        for col, valor in enumerate(datos, start=1):
            c = ws.cell(row=inicio + i, column=col, value=valor)
            c.font = Font(name=FUENTE, size=10)
            c.border = Border(bottom=FINA)
            if formatos[col - 1]:
                c.number_format = formatos[col - 1]
                c.alignment = Alignment(horizontal="right")
    fin = inicio + len(filas) - 1
    if not filas:
        ws.cell(row=inicio, column=1, value="Sin datos en el periodo").font = Font(name=FUENTE, italic=True, color="777777")
    elif total:
        fila_total = fin + 1
        for col, valor in enumerate(total, start=1):
            letra = get_column_letter(col)
            if valor == "suma":
                valor = f"=SUM({letra}{inicio}:{letra}{fin})"
            c = ws.cell(row=fila_total, column=col, value=valor)
            c.font = Font(name=FUENTE, bold=True, size=10)
            c.border = Border(top=GRUESA, bottom=GRUESA)
            if formatos[col - 1]:
                c.number_format = formatos[col - 1]
                c.alignment = Alignment(horizontal="right")
    for col, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def generar_excel(reporte):
    """Recibe el dict de calcular_reporte() y devuelve el .xlsx en memoria (BytesIO)."""
    periodo, r = reporte["periodo"], reporte["resumen"]
    wb = Workbook()

    # ─── Resumen ───
    ws = wb.active
    ws.title = "Resumen"
    fila = _encabezado(ws, "Reporte de mantenimiento vehicular", periodo, 2)
    indicadores = [
        ("Servicios ingresados", r["servicios"], "0"),
        ("Servicios terminados", r["terminados"], "0"),
        ("Siguen en taller", r["en_taller"], "0"),
        ("Gasto total", r["gasto_total"], MONEDA),
        ("Gasto promedio por servicio", r["gasto_promedio"], MONEDA),
        ("Tiempo promedio en taller (días)", r["dias_promedio_taller"], "0.0"),
    ]
    if r.get("sin_costo"):
        indicadores.append(("Terminados sin costo capturado", r["sin_costo"], "0"))
    _tabla(ws, fila, ["Indicador", "Valor"], [(e, v) for e, v, _ in indicadores], [None, None], [38, 18])
    for i, (_, _, formato) in enumerate(indicadores):
        c = ws.cell(row=fila + 1 + i, column=2)
        c.number_format = formato
        c.alignment = Alignment(horizontal="right")
        c.font = Font(name=FUENTE, size=10, bold=True)
    emitido = utcnow() - timedelta(hours=6)  # hora del centro de México
    ws.cell(row=fila + len(indicadores) + 2, column=1,
            value=f"Fecha de emisión: {emitido.strftime('%d/%m/%Y %H:%M')}").font = Font(name=FUENTE, size=9, color="777777")

    # ─── Tablas ───
    def hoja(nombre, encabezados, filas, formatos, anchos, total=None):
        ws = wb.create_sheet(nombre)
        fila = _encabezado(ws, nombre, periodo, len(encabezados))
        _tabla(ws, fila, encabezados, filas, formatos, anchos, total)

    hoja(
        "Gasto por mes", ["Mes", "Servicios", "Terminados", "Gasto"],
        [(_mes(m["mes"]), m["servicios"], m["terminados"], m["gasto"]) for m in reporte["por_mes"]],
        [None, "0", "0", MONEDA], [16, 12, 12, 16], ["Total", "suma", "suma", "suma"],
    )

    def con_porcentaje(lista, clave, etiqueta=lambda x: x):
        total = sum(x["gasto"] for x in lista)
        return [(etiqueta(x[clave]), x["servicios"], x["gasto"], (x["gasto"] / total) if total else None) for x in lista]

    hoja(
        "Gasto por área", ["Área", "Servicios", "Gasto", "% del gasto"],
        con_porcentaje(reporte["por_area"], "area"),
        [None, "0", MONEDA, PORCENTAJE], [40, 12, 16, 13], ["Total", "suma", "suma", "suma"],
    )
    hoja(
        "Gasto por tipo", ["Tipo de servicio", "Servicios", "Gasto", "% del gasto"],
        con_porcentaje(reporte["por_tipo"], "tipo", lambda t: TIPOS.get(t, t)),
        [None, "0", MONEDA, PORCENTAJE], [22, 12, 16, 13], ["Total", "suma", "suma", "suma"],
    )

    def vehiculos(lista):
        return [(i, v["nombre"], v["no_inventario"], v["area"] or "—", v["servicios"], v["gasto"])
                for i, v in enumerate(lista, start=1)]

    columnas = ["#", "Vehículo", "No. inventario", "Área", "Servicios", "Gasto"]
    formatos = ["0", None, None, None, "0", MONEDA]
    anchos = [5, 34, 22, 30, 11, 16]
    hoja("Mayor gasto", columnas, vehiculos(reporte["top_costo"]), formatos, anchos)
    hoja("Más ingresos al taller", columnas, vehiculos(reporte["top_ingresos"]), formatos, anchos)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToHeight = 0

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida
